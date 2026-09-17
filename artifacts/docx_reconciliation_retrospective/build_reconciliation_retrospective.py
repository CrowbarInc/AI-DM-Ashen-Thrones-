from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Master Mandalcio\Documents\Tabletop Gaming\AI Dungeon Master\ashen_thrones_ai_gm")
REFERENCE = ROOT / "history" / "development_logs" / "AI_DM_Project_Narrative_Retrospective.docx"
OUTPUT = ROOT / "history" / "development_logs" / "AI_DM_Project_Narrative_Retrospective_Reconciliation_Era.docx"


TITLE = "The Story of the AI-DM Project"
SUBTITLE = "Reconciliation Era: June-July 2026"


SECTIONS = [
    (
        "Looking Back",
        [
            "The previous chapter ended with the project standing on firmer ground than anyone could reasonably have expected a few months earlier. What began as a question - could an artificial intelligence run a tabletop roleplaying game in a way that felt like a real Game Master? - had become something much more substantial.",
            "By late June, the AI-DM could remember, respond, preserve continuity, and catch many of its own mistakes. It had systems for replaying behavior, tracing failures, protecting important outcomes, and separating the game world's truth from the words used to describe that truth. The project was no longer just proving that an AI Game Master could exist.",
            "It was beginning to prove that such a system could be governed.",
            "That was an important milestone, but it also created a quieter and more difficult problem. The more capable the project became, the harder it became to hold the whole thing in one person's head. There were now layers of game logic, narration, prompts, fallback behavior, validation, replay evidence, persistence, diagnostics, and governance. They worked together, but the shape of that cooperation was no longer obvious at a glance.",
            "So the next chapter did not begin with a flashy new feature. It began with a pause.",
        ],
    ),
    (
        "A Different Kind of Progress",
        [
            "From the outside, the Reconciliation Era could look like a slowdown. Fewer player-facing features appeared. There was no dramatic new combat system, no sweeping world simulation, no polished user interface reveal. Instead, the project spent weeks producing maps, inventories, closeouts, registries, doctrine, and workflow guides.",
            "That kind of work can sound dry until you understand why it mattered.",
            "A software project can reach a point where adding more capability is not the hardest task. The hardest task is making sure the next addition does not quietly damage something that already works. This is especially true for an AI Game Master, because so many parts of the experience depend on trust. If the game says a clue was found, the world needs to remember it. If a character made a promise, the system needs to carry that forward. If the AI improvises, the engine needs to know what is actual truth and what is only narration.",
            "The project had already learned one major lesson: the AI should not be the source of truth. The engine decides what happens; the AI describes it. Reconciliation asked the next question: do we truly understand the engine well enough to grow it?",
            "That question was not theoretical. Feature expansion was about to become more ambitious. Future work might include deeper investigation systems, richer gameplay loops, alternate rulesets, more AI providers, better tools, more visible provenance, and eventually a more polished product experience. Each of those possibilities would touch many systems at once.",
            "Before building outward, the team needed to know what could safely bear weight.",
        ],
    ),
    (
        "Mapping the Unknown",
        [
            "Campaign 1 began with a deceptively simple goal: determine what architecture had actually emerged from Foundation Stabilization.",
            "That word, actually, mattered. The project did not start by assuming the code was broken or that a better design existed somewhere in imagination. It started from evidence. The system had been built organically over months of solving real problems. Some choices were deliberate. Some were pragmatic. Some were transitional. Some looked messy because they were protecting important behavior that was not yet easy to name.",
            "The first task was to recover the map.",
            "That meant identifying the major subsystems and the responsibilities they had quietly taken on. There was a runtime transaction spine that carried a turn from player input through game resolution, narration, final packaging, persistence, logs, and response. There were domain modules that owned the game's truth. There was a state authority layer that decided which parts of the system were allowed to read or change which kinds of state. There was CTIR, a compact meaning layer that captured what had happened before the AI tried to describe it. There was prompt construction, model routing, final emission, persistence, replay, evidence, diagnostics, and governance.",
            "Many of these systems were already doing the right job. The trouble was that their collective shape had become difficult to see.",
            "Reconciliation began by turning a working but mentally expensive system into an understandable one. It was less like inventing a new machine and more like opening the hood, labeling every belt and bracket, and discovering that the engine was more coherent than it first appeared.",
        ],
    ),
    (
        "Asking Bigger Questions",
        [
            "The earliest version of the project asked, \"Can this work?\"",
            "Foundation Stabilization asked, \"Can this keep working when things go wrong?\"",
            "The Reconciliation Era asked a more mature question: \"Why does this architecture work, and will it still work years from now?\"",
            "That shift changed the nature of the work. It was no longer enough to know that a system passed tests or produced reasonable output. The team needed to know why responsibility lived where it did. Why did one layer own truth while another owned expression? Why did replay observe behavior instead of controlling it? Why did final emission own last-mile legality but not the game world's facts? Why were some kinds of fallback split across several owners instead of simplified into one neat box?",
            "The answers were often more subtle than expected.",
            "Some apparent complexity turned out to be real structure. Fallback, for example, was not one thing. There could be fallback content, fallback selection, fallback application, fallback provenance, and fallback observation. Collapsing all of that into a single idea would make the system easier to describe, but less honest about how it actually protected the player experience.",
            "This was one of the era's major discoveries: simplification is not the same thing as clarity. A mature architecture is not the one with the fewest boxes. It is the one where every box has a reason to exist, and where contributors can tell which box owns which kind of decision.",
        ],
    ),
    (
        "Finding the Shape of the Engine",
        [
            "Campaigns 2 and 3 turned the recovered map into a working language.",
            "Campaign 2 compared the architecture against the long-term vision. Could it support modular rulesets? Additional AI backends? Better handoff to future contributors? Low-context operation? Feature extensibility? UI and tooling readiness? Player trust and explainability?",
            "The answer was encouraging: the architecture was broadly compatible with the vision. The missing pieces were mostly local contracts and implementation work, not a reason to redesign the foundation.",
            "That mattered because it prevented the project from falling into a common trap. When a system becomes complex, it is tempting to assume the answer is a grand redesign. Sometimes that is true. Here, the evidence pointed elsewhere. The core architecture did not need to be thrown away. It needed to be understood, named, and extended deliberately.",
            "Campaign 3 then tackled the vocabulary problem. The project had accumulated many terms: replay, realization, ownership, authority, provenance, fallback, validators, repair, sanitizers, diagnostics, evidence, governance, compatibility residue, and more. Some were foundational. Some were implementation concepts. Some were transitional. Some were historical. Some intentionally overlapped because they answered different questions at the same boundary.",
            "That work produced a permanent concept map. Runtime truth flows forward. The engine resolves the world. CTIR captures resolved meaning. Prompt adaptation packages approved truth. The AI produces candidate expression. Final emission decides what can safely be shown. Persistence stores the result. Replay and evidence observe. Governance watches the architecture, but does not play the game.",
            "For a non-technical reader, the important point is this: the team stopped treating the codebase as a collection of clever parts and began treating it as a system with a shared grammar.",
        ],
    ),
    (
        "Drawing the Boundaries",
        [
            "Campaign 4 asked where the walls should be.",
            "Boundaries are one of the least glamorous parts of software, but they are also one of the reasons complex systems survive. A boundary says: this part owns the game's truth; this part adapts truth into narration; this part checks whether final text is safe; this part stores the result; this part observes later; this part documents what must remain true.",
            "Without those boundaries, future work becomes guesswork. A developer adding a feature may not know whether to change the domain logic, the prompt, the final output layer, the replay projection, or the documentation. Worse, they may change all of them just to be safe. That kind of broad coordination slows every feature and increases the chance of accidental damage.",
            "Boundary Reconciliation confirmed that the major ownership lines were durable. The runtime transaction spine should remain the organizing sequence for a turn. Domain simulation should own game truth. State authority should guard mutation and read permissions. CTIR should remain the resolved-turn meaning surface. Prompt adaptation should consume truth rather than invent it. Model output should remain expression, not authority. Final emission should remain the last-mile legality and packaging boundary. Replay and evidence should remain downstream observers. Governance should declare and test doctrine without becoming gameplay logic.",
            "That is a lot of architecture in one paragraph, but the human idea is simple: every important decision now had a home.",
            "And just as important, many things were explicitly not homes for certain decisions. The AI was not truth. Replay was not runtime policy. Documentation was not executable behavior. Provenance explained what happened; it did not choose what happened. Compatibility mechanisms preserved old paths; they did not become rival owners.",
            "The system was learning not only what it was, but what it was not.",
        ],
    ),
    (
        "Making Correct Work Easier",
        [
            "By Campaign 5, the central issue had changed again. The architecture was no longer mysterious. The problem was coordination cost.",
            "Coordination cost is the hidden tax paid every time a contributor needs to ask, \"Where should this change go? Which tests prove it? Which documents must be refreshed? Which old compatibility surfaces matter? Which registry names the owner?\" If those answers are scattered, every feature begins with archaeology.",
            "Campaign 5 reduced that tax.",
            "It created clearer entry points for feature-lane verification, compatibility residue, governance refresh, generated-documentation strategy, backend contracts, ruleset contracts, and version/provenance contracts. It added automation around split-owner reporting. It made future extension points visible without prematurely implementing them.",
            "This was still not glamorous work. Most players will never know that a compatibility register exists. They will not think about governance refresh workflows or backend contract registries. But those tools change the day-to-day reality of development. They make it easier for future contributors to do the right thing on the first attempt.",
            "The project did not eliminate all coordination. It should not. Some coordination is the price of preserving clean ownership. The goal was to remove avoidable rediscovery, not to flatten the architecture into something less capable.",
            "In that sense, Campaign 5 was the era's practical hinge. Earlier campaigns answered, \"What is this system?\" Campaign 5 answered, \"How do we work inside it without getting lost?\"",
        ],
    ),
    (
        "Discovering the Chassis",
        [
            "Campaign 6 gave the era its final metaphor: the chassis.",
            "A chassis is not the paint, the seats, the dashboard, or the engine tuning. It is the underlying structure that lets different models be built without starting from nothing each time. Once the chassis is sound, designers can create new versions with confidence because they know what is load-bearing and what is meant to vary.",
            "That was the final question of Reconciliation. Which parts of Ashen Thrones were now stable enough to support years of future growth? Which systems were intended extension points? Which responsibilities should remain internal details? Which future work should count as normal implementation, and which future work should trigger a real architectural review?",
            "The answer was decisive. The long-term chassis is the forward-flowing architecture recovered across the previous campaigns: player intent enters one runtime transaction; domain systems decide what is true; state authority guards who may change what; CTIR captures resolved meaning; prompt adaptation packages truth for expression; backend and model layers produce candidate language; final emission seals the player-facing result; persistence stores it; replay, diagnostics, evidence, and governance observe and protect it afterward.",
            "That chassis is conservative in the best sense of the word. It does not assume every new idea deserves a new abstraction. It asks future work to attach through known owners and contracts. Ordinary gameplay should extend through domain systems. New state should pass through state authority. New narration meaning should extend CTIR carefully. New AI providers should wait for a backend adapter. Alternate rulesets should wait for ruleset identity and capability contracts. Public or hosted tools should get their own facade when product needs justify it.",
            "The message was not, \"Never change the architecture.\" The message was, \"Change it for evidence, not anxiety.\"",
        ],
    ),
    (
        "What Was Actually Built",
        [
            "The Reconciliation Era produced fewer visible features than an ordinary product sprint might. But it built something larger than a feature.",
            "It built architectural understanding: a shared map of how the project works and why its major responsibilities sit where they do.",
            "It built governance: a way to name owners, classify evidence, protect dependency direction, and prevent old compatibility paths from becoming accidental architecture.",
            "It built documentation that is not merely historical, but operationally useful: closeouts, registries, workflow guides, doctrine maps, and verification lanes that help future work begin from clarity.",
            "It built extension confidence. Backend expansion, ruleset support, version/provenance work, diagnostics, UI improvement, content growth, and gameplay systems now have known attachment points. They may still require work, but they no longer require guessing where the work belongs.",
            "It built restraint. The project learned that not every absence is an architectural gap. Sometimes a feature is simply not implemented yet. Sometimes a contract should be written before expansion. Sometimes a future product decision should remain future until the product actually needs it.",
            "Most of all, it built a new relationship between the team and the codebase. The system was no longer a dense accumulation of solved problems. It had become a known machine with named load-bearing parts, known extension surfaces, and explicit rules for future growth.",
        ],
    ),
    (
        "Why This Matters",
        [
            "Invisible work matters when it changes what future work costs.",
            "Before Reconciliation, a new feature risked becoming a broad expedition. A contributor might need to rediscover the runtime flow, the replay boundary, the final-emission rules, the provenance fields, the compatibility implications, and the relevant tests before safely making progress.",
            "After Reconciliation, that same contributor can start from a doctrine. They can ask more precise questions. What domain owns this behavior? What state changes? Does CTIR need a new slice? Does prompt adaptation consume it? Does final emission need to preserve or validate anything? What replay or provenance evidence should observe the result? Which direct-owner tests prove the change?",
            "That does not make hard features easy. It makes them tractable.",
            "It also protects player trust. An AI Game Master must feel imaginative, but it must not feel arbitrary. The player should be able to rely on the world. Behind that feeling is a tremendous amount of invisible discipline: truth before narration, evidence after runtime, provenance as explanation, compatibility handled deliberately, and governance that watches without interfering.",
            "This is why weeks of documentation can be a major milestone. The work did not merely describe the project. It changed the project's ability to continue.",
            "A prototype can survive on momentum. A long-lived system needs memory, doctrine, and restraint. The Reconciliation Era gave the project those things.",
        ],
    ),
    (
        "Looking Forward",
        [
            "The next era should feel different.",
            "Reconciliation was investigative. Its campaigns moved linearly because each question depended on the previous answer. First recover the map. Then compare it with the vision. Then reconcile the concepts. Then confirm the boundaries. Then reduce coordination cost. Then define the chassis.",
            "Product work does not need to move the same way. Once the chassis is stable, future development can become more portfolio-like: gameplay, AI experience, tools, content, user experience, diagnostics, compatibility cleanup, and performance can each advance when they are the highest-leverage next step.",
            "The important change is psychological as much as technical. Future work no longer needs to begin by wondering whether the foundation is real. It can assume the architecture is complete enough to use, while still respecting the review triggers for genuinely structural changes.",
            "That means controlled expansion: not reckless feature accumulation, and not endless architectural self-study. The next stage should maximize player value through carefully scoped implementation. New features should strengthen confidence in the chassis, not compete with it.",
            "The AI-DM project began by asking whether an artificial intelligence could run a tabletop roleplaying game. It then learned that the deeper question was whether such a game could remain consistent, explainable, and maintainable as it grew.",
            "The Reconciliation Era did not finish the game. It did something quieter and perhaps more important.",
            "It made the project ready to grow up on purpose.",
        ],
    ),
]


def set_run_font(run, name="Aptos", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_style_font(style, name, size=None, bold=None, color=None):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = RGBColor(*color)


def clear_body(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_page_break_after(paragraph):
    run = paragraph.add_run()
    run.add_break()
    br = run._element.br_lst[-1]
    br.set(qn("w:type"), "page")


def add_bottom_border(paragraph, color="8A3B32", size="8", space="14"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_paragraph(doc, text, style="Normal", before=0, after=10, line=1.08):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, "Aptos", 11.2, color=(31, 31, 31))
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, "Aptos Display", 17, bold=True, color=(96, 46, 38))
    return p


def add_title_page(doc):
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(132)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    set_run_font(r, "Aptos Display", 30, bold=True, color=(49, 49, 49))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(SUBTITLE)
    set_run_font(r, "Aptos", 15, color=(96, 46, 38))
    add_bottom_border(p)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("A narrative retrospective on understanding, reconciling, and preparing the AI-DM architecture for long-term growth")
    set_run_font(r, "Aptos", 11.5, italic=True, color=(84, 84, 84))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(168)
    r = p.add_run("Following The Story of the AI-DM Project - March-June 2026")
    set_run_font(r, "Aptos", 10, color=(112, 112, 112))
    add_page_break_after(p)


def build():
    doc = Document(REFERENCE)
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1.12)
    section.right_margin = Inches(1.12)

    clear_body(doc)

    set_style_font(doc.styles["Normal"], "Aptos", 11.2, color=(31, 31, 31))
    doc.styles["Normal"].paragraph_format.space_after = Pt(10)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.08

    set_style_font(doc.styles["Title"], "Aptos Display", 30, bold=True, color=(49, 49, 49))
    set_style_font(doc.styles["Heading 1"], "Aptos Display", 17, bold=True, color=(96, 46, 38))
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(20)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True

    add_title_page(doc)

    for heading, paragraphs in SECTIONS:
        add_heading(doc, heading)
        for text in paragraphs:
            add_paragraph(doc, text)

    for section in doc.sections:
        footer = section.footer
        if footer.paragraphs:
            p = footer.paragraphs[0]
        else:
            p = footer.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("AI-DM Project Narrative Retrospective")
        set_run_font(r, "Aptos", 8.5, color=(130, 130, 130))

    core = doc.core_properties
    core.title = "The Story of the AI-DM Project - Reconciliation Era"
    core.subject = "Narrative retrospective covering the Reconciliation Era, June-July 2026"
    core.keywords = "AI-DM, Ashen Thrones, Reconciliation Era, retrospective, architecture"
    core.comments = "Generated as the next chapter in the AI-DM narrative retrospective series."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
